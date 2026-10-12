# Copyright (c) 2026, ErenAydin - GMP DMS Module
# License: MIT
"""Force one font family across everything the DMS generates.

The problem this solves: the body font of a generated document comes from the
*template*, not from us. LibreOffice renders whatever the .docx asks for, so a
site standardising on a Persian font (Vazir) would otherwise have to re-author
every template — and would still be one careless upload away from a document in
Calibri. These functions rewrite the font references in the rendered file just
before conversion, so the configured family wins regardless of what the author
chose.

Persian specifics, which is what makes this worth doing properly:

* OOXML gives each run THREE font slots — ``w:ascii``/``w:hAnsi`` (Latin),
  ``w:eastAsia`` (CJK) and ``w:cs`` (**complex script**). Persian and Arabic are
  complex-script, so they take ``w:cs``. Setting the font in Word's normal font
  box only writes ``w:ascii``, which is the usual reason "I set the font but the
  Persian text didn't change". We write all four.
* Theme attributes (``w:asciiTheme``, ``w:cstheme``, …) take precedence over the
  explicit family, so they are REMOVED rather than rewritten — otherwise the
  theme silently wins and nothing appears to happen.
* Shaping, ligatures and bidi in the document body are LibreOffice's job and it
  does them correctly (HarfBuzz) once the font resolves. We deliberately do NOT
  touch paragraph direction: forcing RTL would wreck Latin paragraphs, and
  correctly-authored Persian already carries its own direction marks.

reportlab, used for the watermark and footer overlay, is the opposite case: it
has no shaping and no bidi at all, so ``shape_rtl`` below must pre-process any
Persian string before it is drawn.

Deliberately free of Frappe imports so it stays unit-testable in isolation,
matching format_renderers.py.
"""

import os
import re
import zipfile
from xml.sax.saxutils import escape, unescape

from .format_renderers import rewrite_package

# Fonts whose glyphs are pictures, not letters. GMP forms routinely draw
# checkboxes as Wingdings characters, and rewriting those to a text font turns
# every ☑ into a stray letter — a silent corruption of the printed record. Kept
# by default; the caller can switch it off for a literal every-font-replaced
# pass.
#
# MS Gothic is deliberately absent although Word draws its checkbox content
# controls in it: those glyphs are real Unicode (☐ U+2610, ☒ U+2612), which
# LibreOffice draws from a fallback font when the enforced one lacks them, so
# they survive. MS Gothic is also a text font, and preserving it would exempt
# every run that merely names it for East Asian text.
SYMBOL_FONTS = frozenset(
    {
        "wingdings",
        "wingdings 2",
        "wingdings 3",
        "webdings",
        "symbol",
        "zapfdingbats",
        "marlett",
        "bookshelf symbol 7",
        "ms outlook",
        "segoe mdl2 assets",
        "segoe fluent icons",
        "segoe ui symbol",
    }
)

# .docx parts that carry font references. Headers and footers are separate parts
# — miss them and the page furniture keeps the old font while the body changes.
_DOCX_FONT_PARTS = re.compile(
    r"^word/("
    r"document\d*\.xml"
    r"|styles\.xml"
    r"|header\d*\.xml"
    r"|footer\d*\.xml"
    r"|footnotes\.xml"
    r"|endnotes\.xml"
    r"|numbering\.xml"
    r"|comments\.xml"
    r"|theme/theme\d*\.xml"
    r")$"
)

_RFONTS = re.compile(r"<w:rFonts\b([^>]*?)(/?)>")
_ATTR = re.compile(r'(\w+:\w+)\s*=\s*"([^"]*)"')
# DrawingML theme typefaces: <a:latin typeface="..."/>, <a:cs .../>, <a:ea .../>
_THEME_TYPEFACE = re.compile(r'(<a:(?:latin|cs|ea)\b[^>]*?typeface=")([^"]*)(")')
_DOC_DEFAULTS = re.compile(r"<w:docDefaults\b.*?</w:docDefaults>", re.S)

_FONT_SLOTS = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")
_THEME_SLOTS = ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme")


def is_symbol_font(name):
    return (name or "").strip().lower() in SYMBOL_FONTS


def _xml_attr(value):
    """``value`` escaped for a double-quoted XML attribute. The family is a free
    text setting: written raw, a name containing & or " would leave every font
    part malformed — no exception here, just a package soffice cannot open."""
    return escape(value, {'"': "&quot;"})


def _from_xml_attr(value):
    return unescape(value, {"&quot;": '"', "&apos;": "'"}) if value is not None else None


def _rewrite_rfonts(match, family_attr, preserve_symbols, root=False):
    attrs = dict(_ATTR.findall(match.group(1)))
    self_closing = match.group(2)

    if preserve_symbols and any(is_symbol_font(attrs.get(slot)) for slot in _FONT_SLOTS):
        return match.group(0)
    # An element naming no font at all — Word writes <w:rFonts w:hint="cs"/> on
    # runs — only picks a slot; the fonts come from the style chain, which is
    # rewritten where it names them and preserved where that is a symbol font.
    # Expanding it into four explicit slots would override the chain, so a
    # checkbox whose Wingdings comes from a character style would print as a
    # stray letter. docDefaults (``root``) has nothing above it to inherit
    # from, so it is always set.
    if not root and not any(slot in attrs for slot in _FONT_SLOTS + _THEME_SLOTS):
        return match.group(0)

    for slot in _FONT_SLOTS:
        attrs[slot] = family_attr
    # Theme references outrank the explicit family, so drop them entirely.
    for slot in _THEME_SLOTS:
        attrs.pop(slot, None)

    rendered = " ".join(f'{k}="{v}"' for k, v in attrs.items())
    return f"<w:rFonts {rendered}{self_closing}>"


def _typeface_rewriter(family_attr, preserve_symbols):
    """``re.sub`` callback for a ``(prefix, name, suffix)`` typeface match —
    shared by the DrawingML theme and the Visio FaceNames table."""

    def rewrite(match):
        if preserve_symbols and is_symbol_font(match.group(2)):
            return match.group(0)
        return f"{match.group(1)}{family_attr}{match.group(3)}"

    return rewrite


def enforce_docx_font(path, family, preserve_symbols=True):
    """Rewrite every font reference in the .docx at ``path`` to ``family``.

    Covers the body, styles (including ``docDefaults``, which is what runs with
    no explicit font inherit), headers, footers, footnotes, endnotes, numbering
    and the theme. Rewritten in place via rewrite_package, which leaves the
    original untouched if the pass fails; parts without font references are
    copied byte-for-byte.

    Returns the number of font references visited, so a caller can log that
    the pass actually did something.
    """
    if not family:
        return 0

    family_attr = _xml_attr(family)
    theme_rewriter = _typeface_rewriter(family_attr, preserve_symbols)
    rewrites = 0

    def rewrite_rfonts(xml, root=False):
        nonlocal rewrites
        xml, n = _RFONTS.subn(
            lambda m: _rewrite_rfonts(m, family_attr, preserve_symbols, root), xml
        )
        rewrites += n
        return xml

    def rewrite_part(name, data):
        nonlocal rewrites
        if not _DOCX_FONT_PARTS.match(name):
            return data
        xml = data.decode("utf-8")
        defaults = _DOC_DEFAULTS.search(xml)
        if defaults:
            xml = (
                rewrite_rfonts(xml[: defaults.start()])
                + rewrite_rfonts(defaults.group(0), root=True)
                + rewrite_rfonts(xml[defaults.end() :])
            )
        else:
            xml = rewrite_rfonts(xml)
        xml, n = _THEME_TYPEFACE.subn(theme_rewriter, xml)
        rewrites += n
        return xml.encode("utf-8")

    rewrite_package(path, path, rewrite_part)
    return rewrites


def enforce_xlsx_font(workbook, family, preserve_symbols=True):
    """Set ``family`` on every populated cell of an openpyxl workbook.

    Applied to cells rather than the workbook's shared style table, which is a
    private openpyxl structure whose shape changes between versions. That means
    a font set on an empty, never-touched cell is not rewritten — harmless,
    since an empty cell prints nothing.

    Returns the number of cells restyled.
    """
    if not family:
        return 0

    from copy import copy

    changed = 0
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                font = cell.font
                if preserve_symbols and is_symbol_font(getattr(font, "name", None)):
                    continue
                if getattr(font, "name", None) == family:
                    continue
                new_font = copy(font)
                new_font.name = family
                cell.font = new_font
                changed += 1
    return changed


# Visio keeps a table of face names that shapes reference by index, so renaming
# the faces themselves redirects every piece of text on the page.
_VSDX_FACENAME = re.compile(r'(<FaceName\b[^>]*?\bName=")([^"]*)(")')


def enforce_vsdx_font(source_path, out_path, family, preserve_symbols=True):
    """Best-effort font enforcement for a .vsdx, by rewriting the FaceNames
    table that shapes reference.

    Unverified against libvisio's PDF export, like the rest of the Visio image
    path in this module — Visio support in LibreOffice is thin. A part that
    does not decode is copied through unchanged rather than failing the
    render. ``out_path`` may equal ``source_path``. Returns the number of face
    names rewritten.
    """
    if not family:
        return 0

    rewriter = _typeface_rewriter(_xml_attr(family), preserve_symbols)
    rewrites = 0

    def rewrite_part(name, data):
        nonlocal rewrites
        if not (name.startswith("visio/") and name.endswith(".xml")):
            return data
        try:
            xml = data.decode("utf-8")
        except UnicodeDecodeError:
            return data
        xml, n = _VSDX_FACENAME.subn(rewriter, xml)
        rewrites += n
        return xml.encode("utf-8")

    rewrite_package(source_path, out_path, rewrite_part)
    return rewrites


# --------------------------------------------------------------------------- #
#  Reading an enforced font back out of a rendered file                       #
# --------------------------------------------------------------------------- #


def detect_enforced_font(path):
    """Return the family an enforcement pass left in the rendered file at
    ``path``, or None when the file does not carry that signature.

    The signature is uniformity: after a pass every font reference that is not
    a preserved symbol font names one family — in a .docx all four slots of
    every ``w:rFonts`` that names a font (hint-only ones are left alone), with
    no theme attribute left, plus every theme typeface
    (including the empty ``<a:ea typeface=""/>`` a stock theme always has).
    An authored .docx essentially never looks like that. A sheet or drawing
    can simply be uniform in, say, Calibri, so callers must compare the result
    with the family that could have been enforced rather than trust it alone.

    Used to recover the font of documents approved before the font was
    recorded on the document itself (patches/v2_7_1/pin_render_font.py).
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        families = _docx_font_families(path)
    elif ext == ".xlsx":
        families = _xlsx_font_families(path)
    elif ext == ".vsdx":
        families = _vsdx_font_families(path)
    else:
        return None
    if families is None or len(families) != 1:
        return None
    return families.pop() or None


def _docx_font_families(path):
    with zipfile.ZipFile(path) as z:
        parts = [
            z.read(name).decode("utf-8", errors="replace")
            for name in z.namelist()
            if _DOCX_FONT_PARTS.match(name)
        ]
    families = set()
    for xml in parts:
        for match in _RFONTS.finditer(xml):
            attrs = dict(_ATTR.findall(match.group(1)))
            if any(is_symbol_font(attrs.get(slot)) for slot in _FONT_SLOTS):
                continue
            if any(slot in attrs for slot in _THEME_SLOTS):
                return None  # a pass removes every theme attribute
            if not any(slot in attrs for slot in _FONT_SLOTS):
                continue  # names no font (w:hint only); a pass leaves it alone
            # A missing slot adds None, which breaks the uniformity on purpose.
            families.update(_from_xml_attr(attrs.get(slot)) for slot in _FONT_SLOTS)
        for match in _THEME_TYPEFACE.finditer(xml):
            if not is_symbol_font(match.group(2)):
                families.add(_from_xml_attr(match.group(2)))
    return families


def _xlsx_font_families(path):
    from openpyxl import load_workbook

    families = set()
    for ws in load_workbook(path).worksheets:
        for row in ws.iter_rows():
            for cell in row:
                name = getattr(cell.font, "name", None)
                if cell.value is not None and not is_symbol_font(name):
                    families.add(name)
    return families


def _vsdx_font_families(path):
    families = set()
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not (name.startswith("visio/") and name.endswith(".xml")):
                continue
            try:
                xml = z.read(name).decode("utf-8")
            except UnicodeDecodeError:
                continue
            families.update(
                _from_xml_attr(m.group(2))
                for m in _VSDX_FACENAME.finditer(xml)
                if not is_symbol_font(m.group(2))
            )
    return families


# --------------------------------------------------------------------------- #
#  RTL text shaping for reportlab                                             #
# --------------------------------------------------------------------------- #

_ARABIC_RANGE = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")


def has_rtl(text):
    """True when the string contains Arabic-script (incl. Persian) characters."""
    return bool(text) and bool(_ARABIC_RANGE.search(text))


def shape_rtl(text):
    """Return ``text`` ready to be drawn by a renderer with no text shaping.

    reportlab draws glyphs left to right, exactly as they appear in the string,
    and applies no Arabic joining — so Persian handed to it raw comes out as
    disconnected, mirrored letterforms. This applies the two missing steps:
    contextual joining (arabic_reshaper) and the Unicode bidi algorithm
    (python-bidi), producing a visually-ordered string of presentation forms.

    Latin text is returned untouched, so it is safe to call on everything. If
    either library is unavailable the original string is returned rather than
    raising: a watermark is not worth failing a PDF over, and Latin — the only
    thing that rendered before this existed — is unaffected either way.
    """
    if not has_rtl(text):
        return text
    try:
        import arabic_reshaper

        try:
            from bidi import get_display
        except ImportError:  # python-bidi < 0.6
            from bidi.algorithm import get_display

        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def font_covers(font_name, text):
    """True when the reportlab font ``font_name`` has a glyph for every
    character of ``text``.

    reportlab does not complain about a missing glyph: it draws glyph 0, which
    is an empty box or nothing at all. A Persian-only build of a font (Vazir's
    "WOL" variants, many of the B-fonts) has no Latin glyphs, so the English
    status words on a controlled PDF would print blank — an obsolete copy with
    no readable OBSOLETE on it. Callers check first and fall back to Helvetica.

    Built-in Type 1 fonts have no character map to consult and are reported as
    covering everything. A font that cannot be looked up is reported as NOT
    covering, so the caller falls back to Helvetica, which always draws Latin.
    """
    try:
        from reportlab.pdfbase import pdfmetrics

        font = pdfmetrics.getFont(font_name)
    except Exception:
        return False
    cmap = getattr(getattr(font, "face", None), "charToGlyph", None)
    if cmap is None:
        return True
    return all(ord(ch) in cmap for ch in text)
